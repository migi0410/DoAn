from pptx import Presentation
import re

md_path = r"C:\Users\Admin\.gemini\antigravity\brain\f25cabe5-dc3c-49d8-a24d-00d155d270e2\Thesis_Defense_Presentation.md"
out_path = r"C:\Users\Admin\OneDrive\DoAn\Thesis_Presentation.pptx"

with open(md_path, 'r', encoding='utf-8') as f:
    content = f.read()

slides_data = content.split('---')
prs = Presentation()

for slide_text in slides_data:
    slide_text = slide_text.strip()
    if not slide_text:
        continue

    title_match = re.search(r'\*\*Title:\*\*\s*(.*)', slide_text)
    title = title_match.group(1).strip() if title_match else "Slide"

    if "Subtitle:" in slide_text:
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        subtitle_match = re.search(r'\*\*Subtitle:\*\*\s*(.*)', slide_text)
        if subtitle_match and slide.placeholders[1]:
            slide.placeholders[1].text = subtitle_match.group(1).strip()
    else:
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title

        bullets = []
        for line in slide_text.split('\n'):
            line = line.strip()
            if line.startswith('- '):
                bullets.append((0, line[2:]))
            elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. '):
                bullets.append((0, line))
            elif line.startswith('* '):
                bullets.append((0, line[2:]))
            elif line.startswith('-'):
                # Catch some formatting
                pass

        if bullets and slide.placeholders[1]:
            tf = slide.placeholders[1].text_frame
            tf.text = bullets[0][1].replace('**', '')
            for level, text in bullets[1:]:
                p = tf.add_paragraph()
                p.text = text.replace('**', '').replace('*', '')
                p.level = level

    notes_match = re.search(r'\*Speaker Notes:\*\s*"(.*?)"', slide_text, re.DOTALL)
    if notes_match:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes_match.group(1).replace('\n', ' ').strip()

prs.save(out_path)
print("SUCCESS")
